"""Builders for realistic test documents (PDF, DOCX, EPUB)."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

LOREM = (
    "The scheduler decides which process runs next on the processor. "
    "A good policy balances throughput, fairness and response time. "
    "Dr. Smith showed that e.g. round robin scheduling gives every task a turn. "
    "The cost of a context switch is small but not zero, so quanta cannot be too short. "
)


def make_pdf(pages: int = 6, *, bookmarks: bool = True, header: str = "Operating Systems Notes") -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    for page in range(1, pages + 1):
        c.setFont("Helvetica", 9)
        c.drawString(72, height - 40, header)
        c.drawString(width / 2, 30, str(page))
        c.setFont("Helvetica-Bold", 16)
        title = f"Chapter {page}: Topic number {page}"
        c.drawString(72, height - 90, title)
        if bookmarks:
            key = f"p{page}"
            c.bookmarkPage(key)
            c.addOutlineEntry(title, key, level=0)
        c.setFont("Helvetica", 11)
        y = height - 130
        for _ in range(4):
            text = (LOREM * 2).strip()
            words = text.split()
            line = ""
            for w in words:
                if c.stringWidth(line + " " + w, "Helvetica", 11) > width - 144:
                    c.drawString(72, y, line.strip())
                    y -= 15
                    line = w
                else:
                    line += " " + w
            if line:
                c.drawString(72, y, line.strip())
                y -= 15
            y -= 12
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Biology Revision Guide", level=0)
    doc.add_heading("Cells", level=1)
    doc.add_paragraph("The cell is the basic unit of life. " * 12)
    doc.add_paragraph("Mitochondria produce energy.", style="List Bullet")
    doc.add_heading("Genetics", level=1)
    doc.add_paragraph("DNA carries hereditary information. " * 12)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        zf.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>The Test Novel</dc:title></metadata>'
            '<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="c2" href="c2.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
        )
        for i in (1, 2):
            zf.writestr(
                f"OEBPS/c{i}.xhtml",
                f'<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter {i}</h1>'
                f"<p>{'It was a quiet evening in the village. ' * 8}</p>"
                f"<p>{'The river ran beside the old mill. ' * 8}</p></body></html>",
            )
    return buf.getvalue()


_ARTICLE_PARAS = "".join(
    f"<p>{'Sleep is essential for memory, mood and metabolism, and adults need seven to nine hours. ' * 3}</p>"
    for _ in range(6)
)
ARTICLE_HTML = (
    "<!doctype html><html><head><title>Why Sleep Matters | Health Blog</title></head><body>"
    '<nav><a href="/">Home</a><a href="/about">About</a></nav>'
    f"<article><h1>Why Sleep Matters</h1>{_ARTICLE_PARAS}</article>"
    "<footer>Copyright 2026 Health Blog. All rights reserved.</footer></body></html>"
)


def make_scanned_pdf(pages: int = 2, *, lines: list[str] | None = None) -> bytes:
    """A PDF made only of page pictures (no text layer), like a scan."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader

    lines = lines or [
        "Chapter 3: Process Scheduling",
        "",
        "The scheduler decides which process runs next on the processor.",
        "A good policy balances throughput, fairness and response time.",
        "Round robin gives every task a fixed time slice called a quantum.",
        "If the quantum is too short the system wastes time on context switches.",
    ]
    font = ImageFont.load_default(size=30)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    for _ in range(pages):
        img = Image.new("RGB", (1240, 1754), "white")
        draw = ImageDraw.Draw(img)
        y = 140
        for line in lines:
            draw.text((110, y), line, fill="black", font=font)
            y += 62
        c.drawImage(ImageReader(img), 0, 0, width=width, height=height)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_blank_pdf(pages: int) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for _ in range(pages):
        c.showPage()
    c.save()
    return buf.getvalue()
